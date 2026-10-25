from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ── Colour palette ────────────────────────────────────────────────────────────
_DISH_FILL    = PatternFill(start_color="D6E4F7", end_color="D6E4F7", fill_type="solid")
_TOTAL_FILL   = PatternFill(start_color="D5E8D4", end_color="D5E8D4", fill_type="solid")
_ROUNDED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
_HEADER_FILL  = PatternFill(start_color="404040", end_color="404040", fill_type="solid")

_BOLD       = Font(bold=True)
_WHITE_BOLD = Font(bold=True, color="FFFFFF")

_COLS     = ["ingredient", "uom", "type", "exact_qty", "pack_qty", "qty_was_rounded"]
_HEADERS  = ["Ingredient", "UOM", "Type", "Exact Qty", "Pack Qty", "↑"]
_NUM_COLS = {4, 5}
_FLAG_COL = 6

_SVC_TYPE = "packaging / serviceware"   # lowercase for comparison


def _is_serviceware(type_val) -> bool:
    return str(type_val).strip().lower() == _SVC_TYPE


def _filter(df: pd.DataFrame, serviceware: bool) -> pd.DataFrame:
    mask = df["type"].apply(_is_serviceware)
    return df[mask] if serviceware else df[~mask]


# ── Sheet-writing helpers ─────────────────────────────────────────────────────

def _write_col_headers(ws, row: int) -> None:
    for col, header in enumerate(_HEADERS, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = _WHITE_BOLD
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center" if col > 1 else "left")


def _write_section_label(ws, row: int, label: str, n_cols: int, fill) -> None:
    for col in range(1, n_cols + 1):
        ws.cell(row=row, column=col).fill = fill
    cell = ws.cell(row=row, column=1, value=label)
    cell.font = _BOLD
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 16


def _write_data_row(ws, row: int, record: dict) -> None:
    is_rounded = bool(record.get("qty_was_rounded"))
    fill = _ROUNDED_FILL if is_rounded else None
    for col, key in enumerate(_COLS, 1):
        val = record.get(key)
        if key == "qty_was_rounded":
            val = "⬆" if is_rounded else ""
        cell = ws.cell(row=row, column=col, value=val)
        if fill:
            cell.fill = fill
        if col in _NUM_COLS and isinstance(val, (int, float)):
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = "0.00" if isinstance(val, float) and val != int(val) else "0"
        if col == _FLAG_COL and val == "⬆":
            cell.font = Font(bold=True, color="C55A11")
            cell.alignment = Alignment(horizontal="center")


def _autofit(ws) -> None:
    min_widths = [48, 8, 30, 11, 10, 4]
    for col_idx, min_w in enumerate(min_widths, 1):
        letter = get_column_letter(col_idx)
        max_len = max(
            (len(str(c.value)) for c in ws[letter] if c.value is not None),
            default=min_w,
        )
        ws.column_dimensions[letter].width = max(min_w, min(max_len + 3, 60))


def _populate_sheet(ws, sections, total_df, serviceware: bool) -> None:
    """Write one sheet — either Food or Serviceware — with the same grouped layout."""
    n_cols = len(_COLS)
    current_row = 1

    _write_col_headers(ws, current_row)
    ws.freeze_panes = f"A{current_row + 1}"
    current_row += 1

    for sec in sections:
        rows = _filter(sec.df, serviceware)
        if rows.empty:
            continue                       # dish has nothing for this sheet

        mult_str = (
            f"{sec.multiplier:.4g}"
            if sec.multiplier != int(sec.multiplier)
            else str(int(sec.multiplier))
        )
        label = f"{sec.invoice_item}  ×{mult_str}  →  {sec.bom_recipe}"
        _write_section_label(ws, current_row, label, n_cols, _DISH_FILL)
        current_row += 1

        for _, row in rows.iterrows():
            _write_data_row(ws, current_row, row.to_dict())
            current_row += 1

    total_rows = _filter(total_df, serviceware)
    if not total_rows.empty:
        _write_section_label(ws, current_row, "GRAND TOTAL", n_cols, _TOTAL_FILL)
        current_row += 1
        for _, row in total_rows.iterrows():
            _write_data_row(ws, current_row, row.to_dict())
            current_row += 1

    _autofit(ws)


# ── Public entry point ────────────────────────────────────────────────────────

def export_to_excel(result) -> bytes:
    """Write a two-sheet workbook: Food and Serviceware.

    Food sheet       — Purchased Ingredient + Prepared Item / Sub-recipe,
                       grouped by dish with a Grand Total section.
    Serviceware sheet — Packaging / Serviceware rows only, same layout.
                        Boss uses this for the utensils packing list.
    """
    from utils.packing_generator import PackingResult

    if isinstance(result, PackingResult):
        sections = result.sections
        total_df = result.df
    else:
        sections = []
        total_df = result

    wb = Workbook()

    ws_food = wb.active
    ws_food.title = "Food"
    _populate_sheet(ws_food, sections, total_df, serviceware=False)

    ws_svc = wb.create_sheet("Serviceware")
    _populate_sheet(ws_svc, sections, total_df, serviceware=True)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
